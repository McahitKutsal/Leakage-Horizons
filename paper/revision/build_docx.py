"""Assemble the section files into one Word manuscript.

Reads sections/*.md in order, writes manuscript.md (for diffing and word counts)
and manuscript.docx (the submission file).

Markdown subset understood here:
  #/##/###/####   title and headings
  **bold** *italic* `code`
  $inline math$ and $$display math$$   rendered as Unicode with real sub/superscripts
  | tables |                           rendered booktabs-style
  - bullets, 1. numbered
  [FIGURE n]                           embeds figures/figure-n.png
  **Table N.** / **Figure N.**         rendered as a caption paragraph
"""
import re
import pathlib
import sys

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

HERE = pathlib.Path(__file__).parent
SECTIONS = [
    "00-front.md", "01-introduction.md", "02-what-this-is.md", "03-related.md",
    "04-formulation.md", "05-suite.md", "06-design.md",
    "07a-results-synthetic.md", "07b-results-field.md",
    "07c-results-scope-task.md", "08-misreports.md", "09-one-date.md",
    "10-limitations.md", "11-conclusion.md", "12-references.md",
    "13-backmatter.md",
]

AUTHORS = [
    ("Mücahit Kutsal", "a,b", "mucahitkutsal@ankara.edu.tr"),
    ("Alpaslan Erdağ", "a,b", "alpaslan.erdag@ankara.edu.tr"),
    ("Uğur Özbek", "b,c", "uozbek@mfa.gov.tr"),
    ("Doğa Koçyiğit", "d", "doga.kocyigit@stud.tu-darmstadt.de"),
    ("Ömür Koçyiğit", "b", "omur.kocyigit@mfa.gov.tr"),
]
AFFILIATIONS = [
    ("a", "Ankara University, Ankara, Türkiye"),
    ("b", "Republic of Türkiye Ministry of Foreign Affairs, Ankara, Türkiye"),
    ("c", "Konya Technical University, Konya, Türkiye"),
    ("d", "Technical University of Darmstadt, Darmstadt, Germany"),
]

BODY_FONT = "Times New Roman"
BODY_SIZE = Pt(12)
LINE_SPACING = 1.5

# ---------------------------------------------------------------- maths

GREEK = {
    "alpha": "α", "beta": "β", "gamma": "γ", "delta": "δ", "Delta": "Δ",
    "epsilon": "ε", "varepsilon": "ε", "eta": "η", "theta": "θ", "kappa": "κ",
    "lambda": "λ", "Lambda": "Λ", "mu": "μ", "nu": "ν", "pi": "π", "Pi": "Π",
    "rho": "ρ", "sigma": "σ", "Sigma": "Σ", "tau": "τ", "phi": "φ", "Phi": "Φ",
    "chi": "χ", "psi": "ψ", "omega": "ω", "Omega": "Ω",
}
SYMBOLS = {
    "star": "⋆", "geq": "≥", "leq": "≤", "neq": "≠", "approx": "≈",
    "in": "∈", "subseteq": "⊆", "subset": "⊂", "cap": "∩", "cup": "∪",
    "ldots": "…", "cdots": "⋯", "times": "×", "pm": "±", "to": "→",
    "mid": "|", "quad": "  ", "qquad": "    ", ",": " ", ";": " ", "!": "",
    "left": "", "right": "", "big": "", "Big": "", "bigg": "", "Bigg": "",
    "displaystyle": "", "text": "", "mathrm": "", "mspace": "",
}

def _replace_frac(s):
    """\\frac{a}{b} -> (a) / (b), with brace matching, so either side may
    itself contain braces (a subscript, say)."""
    def grab(text, start):
        assert text[start] == "{"
        depth, j = 1, start + 1
        while j < len(text) and depth:
            depth += (text[j] == "{") - (text[j] == "}")
            j += 1
        return text[start + 1:j - 1], j

    while True:
        at = s.find("\\frac{")
        if at == -1:
            return s
        num, j = grab(s, at + 5)
        if j >= len(s) or s[j] != "{":
            return s.replace("\\frac", "", 1)
        den, k = grab(s, j)
        s = s[:at] + f"({num}) / ({den})" + s[k:]


def _strip_math_wrappers(s):
    s = s.replace("\\;", " ").replace("\\:", " ").replace("\\ ", " ")
    s = re.sub(r"\\left|\\right", "", s)
    # \text{...} and \mathrm{...} keep their contents upright
    s = re.sub(r"\\(?:text|mathrm|mathbf|textbf)\{([^{}]*)\}", r"⟦\1⟧", s)
    s = _replace_frac(s)
    s = re.sub(r"\\binom\{([^{}]*)\}\{([^{}]*)\}", r"C(\1, \2)", s)
    s = re.sub(r"\\sqrt\{([^{}]*)\}", r"√(\1)", s)
    s = re.sub(r"\\widehat\{([^{}]*)\}", "\\1\u0302", s)
    s = re.sub(r"\\hat\{([^{}]*)\}", "\\1\u0302", s)
    s = re.sub(r"\\min\b", "⟦min⟧", s)
    s = re.sub(r"\\max\b", "⟦max⟧", s)
    s = re.sub(r"\\log\b", "⟦log⟧", s)
    # a trailing "_" or "^" is not a word boundary, so \b would miss \phi_{r}
    for name, glyph in sorted(GREEK.items(), key=lambda kv: -len(kv[0])):
        s = re.sub(r"\\" + name + r"(?![A-Za-z])", glyph, s)
    for name, glyph in sorted(SYMBOLS.items(), key=lambda kv: -len(kv[0])):
        if name.isalpha():
            s = re.sub(r"\\" + name + r"(?![A-Za-z])", glyph, s)
        else:
            s = s.replace("\\" + name, glyph)
    s = s.replace("\\\\", " ")
    s = re.sub(r"\\[a-zA-Z]+", "", s)
    s = re.sub(r"(?<=[\s\w)\]])-(?=[\s(\[])", "−", s)   # binary minus
    return s


def math_runs(src):
    """Turn a LaTeX fragment into [(text, italic, sub, sup)] runs."""
    s = _strip_math_wrappers(src)
    out = []
    i = 0
    while i < len(s):
        ch = s[i]
        if ch in "_^" and i + 1 < len(s):
            kind = "sub" if ch == "_" else "sup"
            i += 1
            if s[i] == "{":
                depth, j = 1, i + 1
                while j < len(s) and depth:
                    depth += (s[j] == "{") - (s[j] == "}")
                    j += 1
                body, i = s[i + 1:j - 1], j
            else:
                body, i = s[i], i + 1
            # nested scripts cannot be nested again in Word, so flatten them
            body = (body.replace("{", "").replace("}", "")
                        .replace("_", "").replace("^", ""))
            for piece, italic in _split_italic(body):
                out.append((piece, italic, kind == "sub", kind == "sup"))
            continue
        if ch in "{}":
            i += 1
            continue
        j = i
        while j < len(s) and s[j] not in "_^{}":
            j += 1
        for piece, italic in _split_italic(s[i:j]):
            out.append((piece, italic, False, False))
        i = j
    return [(t, it, sub, sup) for (t, it, sub, sup) in out if t]


def _split_italic(text):
    """Letters are italic in maths; digits, operators and ⟦upright⟧ spans are not."""
    pieces, buf, buf_italic = [], "", None
    k = 0
    while k < len(text):
        if text[k] == "⟦":
            end = text.find("⟧", k)
            if end == -1:
                end = len(text)
            if buf:
                pieces.append((buf, buf_italic))
                buf, buf_italic = "", None
            pieces.append((text[k + 1:end], False))
            k = end + 1
            continue
        ch = text[k]
        italic = ch.isalpha() and ch not in "⋆★"
        if buf_italic is None or italic == buf_italic:
            buf, buf_italic = buf + ch, italic
        else:
            pieces.append((buf, buf_italic))
            buf, buf_italic = ch, italic
        k += 1
    if buf:
        pieces.append((buf, buf_italic))
    return pieces


# ---------------------------------------------------------------- inline markup

TOKEN = re.compile(r"(\$\$.+?\$\$|\$[^$]+?\$|\*\*.+?\*\*|\*[^*]+?\*|`[^`]+?`)",
                   re.S)


def add_inline(paragraph, text, base_bold=False, base_italic=False):
    for part in TOKEN.split(text):
        if not part:
            continue
        if part.startswith("$$") and part.endswith("$$"):
            _add_math(paragraph, part[2:-2], bold=base_bold)
        elif part.startswith("$") and part.endswith("$") and len(part) > 2:
            _add_math(paragraph, part[1:-1], bold=base_bold)
        elif part.startswith("**") and part.endswith("**"):
            add_inline(paragraph, part[2:-2], True, base_italic)
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            add_inline(paragraph, part[1:-1], base_bold, True)
        elif part.startswith("`") and part.endswith("`"):
            r = _run(paragraph, part[1:-1], bold=base_bold, italic=base_italic)
            r.font.name = "Consolas"
            r.font.size = Pt(10.5)
        else:
            _run(paragraph, part, bold=base_bold, italic=base_italic)


def _add_math(paragraph, src, bold=False):
    for text, italic, sub, sup in math_runs(src):
        r = _run(paragraph, text, bold=bold, italic=italic)
        if sub:
            r.font.subscript = True
        if sup:
            r.font.superscript = True


def _run(paragraph, text, bold=False, italic=False):
    r = paragraph.add_run(text)
    r.bold = bold
    r.italic = italic
    r.font.name = BODY_FONT
    r.font.size = BODY_SIZE
    rpr = r._element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rpr.insert(0, rf)
    for attr in ("w:ascii", "w:hAnsi", "w:cs"):
        rf.set(qn(attr), BODY_FONT)
    return r


# ---------------------------------------------------------------- doc helpers

def para(doc, text="", style=None, align=None, space_after=6, spacing=None,
         bold=False, italic=False, size=None, indent=None):
    p = doc.add_paragraph(style=style)
    pf = p.paragraph_format
    pf.space_after = Pt(space_after)
    pf.space_before = Pt(0)
    pf.line_spacing = LINE_SPACING if spacing is None else spacing
    pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    if align is not None:
        p.alignment = align
    if indent is not None:
        pf.left_indent = Inches(indent)
    if text:
        add_inline(p, text, base_bold=bold, base_italic=italic)
    if size is not None:
        for r in p.runs:
            r.font.size = size
    return p


def set_cell_border(cell, **kwargs):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "bottom", "left", "right"):
        spec = kwargs.get(edge)
        tag = f"w:{edge}"
        el = borders.find(qn(tag))
        if el is None:
            el = OxmlElement(tag)
            borders.append(el)
        if spec is None:
            el.set(qn("w:val"), "nil")
        else:
            el.set(qn("w:val"), "single")
            el.set(qn("w:sz"), str(spec))
            el.set(qn("w:color"), "000000")


def add_table(doc, rows):
    header, body = rows[0], rows[1:]
    t = doc.add_table(rows=len(rows), cols=len(header))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = True
    for j, text in enumerate(header):
        cell = t.cell(0, j)
        cell.text = ""
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.space_before = Pt(3)
        p.paragraph_format.line_spacing = 1.0
        add_inline(p, text, base_bold=True)
        for r in p.runs:
            r.bold = True
            r.font.size = Pt(9.5)
        set_cell_border(cell, top=12, bottom=6, left=None, right=None)
    for i, row in enumerate(body, start=1):
        last = (i == len(body))
        for j in range(len(header)):
            cell = t.cell(i, j)
            cell.text = ""
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.line_spacing = 1.0
            add_inline(p, row[j] if j < len(row) else "")
            for r in p.runs:
                r.font.size = Pt(9.5)
            set_cell_border(cell, top=None, bottom=12 if last else None,
                            left=None, right=None)
    return t


def add_figure(doc, name):
    path = HERE / "figures" / name
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.0
    p.add_run().add_picture(str(path), width=Inches(6.1))


def split_row(line):
    line = line.strip()
    if line.startswith("|"):
        line = line[1:]
    if line.endswith("|"):
        line = line[:-1]
    return [c.strip() for c in line.split("|")]


def is_rule(line):
    return bool(re.fullmatch(r"\|[\s:|-]+\|", line.strip()))


# ---------------------------------------------------------------- build

def build(src="sections", names=None, stem="manuscript"):
    text = []
    for name in (names or SECTIONS):
        text.append((HERE / src / name).read_text(encoding="utf-8").rstrip())
    md = "\n\n".join(text) + "\n"
    (HERE / f"{stem}.md").write_text(md, encoding="utf-8")

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = BODY_FONT
    style.font.size = BODY_SIZE
    style.element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)

    for sec in doc.sections:
        sec.top_margin = sec.bottom_margin = Inches(1.0)
        sec.left_margin = sec.right_margin = Inches(1.0)
        ln = OxmlElement("w:lnNumType")
        ln.set(qn("w:countBy"), "1")
        ln.set(qn("w:restart"), "continuous")
        ln.set(qn("w:distance"), "360")
        sec._sectPr.append(ln)

    lines = md.split("\n")
    i = 0
    fig_counter = {"n": 0}
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        if not stripped:
            i += 1
            continue

        # title page
        if stripped.startswith("# "):
            p = para(doc, stripped[2:], align=WD_ALIGN_PARAGRAPH.CENTER,
                     space_after=14, spacing=1.2)
            for r in p.runs:
                r.bold = True
                r.font.size = Pt(16)
            add_title_block(doc)
            i += 1
            continue

        if stripped.startswith("#### "):
            p = para(doc, stripped[5:], space_after=4, spacing=1.2, bold=True)
            p.paragraph_format.space_before = Pt(10)
            for r in p.runs:
                r.bold = True
                r.font.size = Pt(12)
            i += 1
            continue

        if stripped.startswith("### "):
            p = para(doc, stripped[4:], space_after=5, spacing=1.2, bold=True)
            p.paragraph_format.space_before = Pt(12)
            for r in p.runs:
                r.bold = True
                r.font.size = Pt(12.5)
            i += 1
            continue

        if stripped.startswith("## "):
            p = para(doc, stripped[3:], space_after=7, spacing=1.2, bold=True)
            p.paragraph_format.space_before = Pt(16)
            for r in p.runs:
                r.bold = True
                r.font.size = Pt(14)
            i += 1
            continue

        # figure placeholder
        m = re.fullmatch(r"\[FIGURE (\w+)\]", stripped)
        if m:
            key = m.group(1)
            add_figure(doc, f"figure-{key}.png")
            fig_counter["n"] += 1
            i += 1
            continue

        # table
        if stripped.startswith("|") and i + 1 < len(lines) and is_rule(lines[i + 1]):
            rows = [split_row(stripped)]
            j = i + 2
            while j < len(lines) and lines[j].strip().startswith("|"):
                rows.append(split_row(lines[j]))
                j += 1
            add_table(doc, rows)
            i = j
            continue

        # caption
        if re.match(r"\*\*(Table|Figure) \d+\.\*\*", stripped):
            p = para(doc, stripped, space_after=12, spacing=1.0,
                     align=WD_ALIGN_PARAGRAPH.LEFT)
            p.paragraph_format.space_before = Pt(5)
            for r in p.runs:
                r.font.size = Pt(10)
            i += 1
            continue

        # display maths on its own line
        if stripped.startswith("$$") and stripped.endswith("$$"):
            p = para(doc, stripped, align=WD_ALIGN_PARAGRAPH.CENTER,
                     space_after=10, spacing=1.0)
            p.paragraph_format.space_before = Pt(10)
            i += 1
            continue

        # bullets
        if stripped.startswith("- "):
            para(doc, stripped[2:], style="List Bullet", space_after=4)
            i += 1
            continue

        m = re.match(r"(\d+)\.\s+(.*)", stripped)
        if m:
            para(doc, m.group(2), style="List Number", space_after=4)
            i += 1
            continue

        # ordinary paragraph: gather until blank line
        buf = [stripped]
        j = i + 1
        while j < len(lines) and lines[j].strip() and not _breaks(lines[j]):
            buf.append(lines[j].strip())
            j += 1
        para(doc, " ".join(buf), align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=8)
        i = j

    out = HERE / f"{stem}.docx"
    doc.save(out)
    words = len(re.sub(r"[^\w\s]", " ", md).split())
    for name, _, email in AUTHORS:
        if email.startswith("["):
            print(f"  PLACEHOLDER: {name} has no e-mail on the byline")
    for mark, where in AFFILIATIONS:
        if where.startswith("["):
            print(f"  PLACEHOLDER: affiliation {mark} is not filled in")
    print(f"wrote {out}")
    print(f"wrote {HERE / (stem + '.md')}  ({words:,} tokens incl. tables)")


def _breaks(line):
    s = line.strip()
    return (s.startswith("#") or s.startswith("|") or s.startswith("- ")
            or s.startswith("$$") or s.startswith("[FIGURE")
            or bool(re.match(r"\d+\.\s", s))
            or bool(re.match(r"\*\*(Table|Figure) \d+\.\*\*", s)))


def add_title_block(doc):
    p = para(doc, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=8, spacing=1.2)
    for idx, (name, marks, _) in enumerate(AUTHORS):
        if idx:
            _run(p, ", ")
        _run(p, name)
        _run(p, marks).font.superscript = True
    for mark, where in AFFILIATIONS:
        p = para(doc, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2, spacing=1.0)
        r = _run(p, mark)
        r.font.superscript = True
        r.font.size = Pt(10)
        r = _run(p, " " + where)
        r.font.size = Pt(10)
        r.italic = True
    p = para(doc, "Corresponding author: Mücahit Kutsal, mucahitkutsal@ankara.edu.tr",
             align=WD_ALIGN_PARAGRAPH.CENTER, space_after=16, spacing=1.0)
    for r in p.runs:
        r.font.size = Pt(10)
        r.italic = True


SUPPLEMENT = ["S0-front.md", "S2-robustness.md", "S4-filters.md",
              "S3-controls.md"]

if __name__ == "__main__":
    build()
    build(src="supplement", names=SUPPLEMENT, stem="supplementary")
