"""Dump what actually landed in the .docx, so the maths and the floats can be eyed."""
import sys
import zipfile

from docx import Document

doc = Document("manuscript.docx")

print("PARAGRAPHS:", len(doc.paragraphs), " TABLES:", len(doc.tables))
with zipfile.ZipFile("manuscript.docx") as z:
    media = [n for n in z.namelist() if n.startswith("word/media/")]
print("EMBEDDED IMAGES:", len(media))
for m in sorted(media):
    print("   ", m)

want = sys.argv[1:] or ["leakage horizon", "phi", "widehat", "Delta"]
print("\nFIRST 28 PARAGRAPHS")
for p in doc.paragraphs[:28]:
    if p.text.strip():
        print(f"  [{p.style.name:14s}] {p.text[:120]}")

print("\nPARAGRAPHS CONTAINING MATHS")
shown = 0
for p in doc.paragraphs:
    runs = p.runs
    if any(r.font.subscript or r.font.superscript for r in runs):
        marked = ""
        for r in runs:
            if r.font.subscript:
                marked += f"_({r.text})"
            elif r.font.superscript:
                marked += f"^({r.text})"
            else:
                marked += r.text
        print("  " + marked[:190])
        shown += 1
        if shown >= 14:
            break

print("\nTABLE SHAPES AND HEADERS")
for i, t in enumerate(doc.tables, 1):
    hdr = " | ".join(c.text[:22] for c in t.rows[0].cells)
    print(f"  T{i}: {len(t.rows)}x{len(t.columns)}   {hdr[:140]}")

print("\nANY LEFTOVER LATEX OR MARKDOWN")
bad = []
for p in doc.paragraphs:
    for token in ("\\", "$", "**", "](", "{", "}"):
        if token in p.text:
            bad.append((token, p.text[:120]))
            break
for t in doc.tables:
    for row in t.rows:
        for c in row.cells:
            for token in ("\\", "$", "**", "{", "}"):
                if token in c.text:
                    bad.append((token, "TABLE: " + c.text[:110]))
                    break
if bad:
    for token, text in bad[:25]:
        print(f"  {token!r:6s} {text}")
    print(f"  ... {len(bad)} total")
else:
    print("  none")
